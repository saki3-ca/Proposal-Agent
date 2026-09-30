"""Snapshot -> XLSX and DOCX downloads."""
import os
import re
import tempfile
from typing import Any, Dict

from backend.tor.render import FIELD_TITLES, TITLES, parsed_text


def _base(snap: Dict[str, Any]) -> str:
    name = os.path.splitext(snap["meta"]["filename"])[0]
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", name)[:60] or "tender"


def _out_dir() -> str:
    return tempfile.mkdtemp(prefix="snapshot_")


def to_xlsx(snap: Dict[str, Any]) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    meta = snap["meta"]
    ws.append(["File", meta["filename"]])
    ws.append(["Pages", meta["pages"]])
    ws.append(["Text extraction", meta["extraction_method"]])
    ws.append(["Model(s)", ", ".join(f"{m['provider']}/{m['model']}" for m in meta["models_used"])])
    ws.append(["Analysis date", meta["analysis_date"]])
    for k, v in meta["counts"].items():
        ws.append([k, v])

    def sheet(title, header, rows):
        s = wb.create_sheet(title[:31])
        s.append(header)
        for c in s[1]:
            c.font = Font(bold=True)
        for r in rows:
            s.append(r)
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width = 45
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    sheet("Key facts", ["Field", "Value", "Parsed (by code)", "Page", "Source quote"],
          [[FIELD_TITLES[f["label"]], f["value"], parsed_text(f.get("parsed")), f["page"], f["source_quote"]] for f in snap["key_facts"]])
    sheet("Evaluation", ["Criterion", "Marks/weight (raw)", "Parsed", "Page", "Source quote"],
          [[r["label"], r["value"], parsed_text(r.get("parsed")), r["page"], r["source_quote"]] for r in snap["evaluation"]])
    sheet("Team", ["Position", "Qualification", "Experience (raw)", "Parsed", "Page", "Source quote"],
          [[r["label"], r.get("qualification", ""), r["value"], parsed_text(r.get("parsed")), r["page"], r["source_quote"]] for r in snap["team"]])
    for section in ("eligibility", "scope", "deliverables", "timeline", "documents", "red_flags"):
        sheet(TITLES[section], ["Label", "Page", "Source quote", "Origin"],
              [[r.get("label", ""), r["page"], r["source_quote"], r.get("origin", "llm")] for r in snap[section]])
    sheet("Unmapped - review", ["Page", "Sentence", "Why flagged"],
          [[u["page"], u["sentence"], ", ".join(u["reasons"])] for u in snap["unmapped"]])
    sheet("Dropped (no quote)", ["Section", "Label", "Page claimed", "Reason", "Claimed quote"],
          [[d["section"], d["label"], d["page"], d["reason"], d["source_quote"]] for d in snap["dropped"]])
    path = os.path.join(_out_dir(), f"{_base(snap)}_snapshot.xlsx")
    wb.save(path)
    return path


def to_docx(snap: Dict[str, Any]) -> str:
    from docx import Document

    doc = Document()
    meta = snap["meta"]
    doc.add_heading(f"Tender Snapshot - {meta['filename']}", 0)
    doc.add_paragraph(
        f"Pages: {meta['pages']} | Extraction: {meta['extraction_method']} | Models: "
        + (", ".join(f"{m['provider']}/{m['model']}" for m in meta["models_used"]) or "none")
        + f" | Analysis date: {meta['analysis_date']}")
    if meta["chunk_errors"]:
        doc.add_paragraph("WARNING: some pages failed analysis; snapshot is incomplete.")

    def table(header, rows):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)

    doc.add_heading("Key facts", 1)
    table(["Field", "Value", "Parsed", "Page", "Source quote"],
          [[FIELD_TITLES[f["label"]], f["value"], parsed_text(f.get("parsed")), f"p.{f['page']}", f["source_quote"]] for f in snap["key_facts"]])
    doc.add_heading("Evaluation criteria", 1)
    table(["Criterion", "Marks / weight", "Parsed", "Page", "Source quote"],
          [[r["label"], r["value"], parsed_text(r.get("parsed")), f"p.{r['page']}", r["source_quote"]] for r in snap["evaluation"]])
    doc.add_heading("Team / key personnel requirements", 1)
    table(["Position", "Qualification", "Experience", "Page", "Source quote"],
          [[r["label"], r.get("qualification", ""), r["value"], f"p.{r['page']}", r["source_quote"]] for r in snap["team"]])
    for section in ("eligibility", "scope", "deliverables", "timeline", "documents", "red_flags"):
        doc.add_heading(TITLES[section], 1)
        for r in snap[section]:
            label = f"{r['label']}: " if r.get("label") else ""
            doc.add_paragraph(f"{label}\"{r['source_quote']}\" (p.{r['page']})", style="List Bullet")
    doc.add_heading("Unmapped - review", 1)
    table(["Page", "Sentence", "Why flagged"], [[f"p.{u['page']}", u["sentence"], ", ".join(u["reasons"])] for u in snap["unmapped"]])
    path = os.path.join(_out_dir(), f"{_base(snap)}_snapshot.docx")
    doc.save(path)
    return path
