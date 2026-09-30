"""Page-aware ingestion on top of the existing MarkItDown + OCR pipeline.

PDF  : MarkItDown text, split per page on form-feeds (checked against the real page
       count, else per-page pdfium text). Pages with no text layer are OCR'd page by page.
DOCX : body walked in order (paragraphs + table rows); page numbers come from Word's
       own page-break markers. If the file carries none, everything is page 1 and the
       method string says so.
OCR  : Tesseract (eng+ben) when installed - needed for scanned Bangla - else RapidOCR
       via the existing OcrEngine (good for English, weak on Bangla).
"""
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
from typing import Dict, Tuple

from backend.services.document_processor import DocumentProcessor
from backend.services.ocr_engine import OcrEngine
from backend.tor.textutil import SourceDoc

logger = logging.getLogger("tor.pages")

MIN_PAGE_CHARS = 40  # below this a PDF page is treated as scanned
_processor = None


def _converter():
    global _processor
    if _processor is None:
        _processor = DocumentProcessor()
    return _processor.md_converter


def _ocr_image(pil_image) -> Tuple[str, str]:
    if shutil.which("tesseract"):
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "p.png")
            pil_image.save(path)
            langs = os.getenv("OCR_LANGS", "eng+ben")
            proc = subprocess.run(["tesseract", path, "stdout", "-l", langs, "--psm", "6"],
                                  capture_output=True, text=True, timeout=180)
            if proc.returncode == 0 and proc.stdout.strip():
                return proc.stdout, f"tesseract({langs})"
    import numpy as np
    engine = OcrEngine.get_ocr_instance()
    if engine is None:
        return "", "no-ocr-engine"
    result, _ = engine(np.array(pil_image.convert("RGB")))
    if not result:
        return "", "rapidocr"
    lines = sorted(result, key=lambda it: (round(it[0][0][1] / 15) * 15, it[0][0][0]))
    return "\n".join(it[1] for it in lines if it[1].strip()), "rapidocr"


def _pdf_pages(data: bytes, filename: str) -> Tuple[Dict[int, str], str]:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    n = len(pdf)
    method = "markitdown"
    pages: Dict[int, str] = {}

    text = ""
    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    try:
        tmp.write(data)
        tmp.close()
        conv = _converter()
        if conv is not None:
            text = (conv.convert(tmp.name).text_content or "")
    except Exception as exc:
        logger.warning("MarkItDown failed on %s: %s", filename, exc)
    finally:
        os.unlink(tmp.name)

    parts = text.split("\f")
    if len(parts) == n:
        pages = {i + 1: p for i, p in enumerate(parts)}
    else:
        method = "pdfium-text (MarkItDown page split did not match page count)"
        for i in range(n):
            pages[i + 1] = pdf[i].get_textpage().get_text_range()

    ocr_used = set()
    for i in range(n):
        if len(re.sub(r"\s", "", pages[i + 1])) < MIN_PAGE_CHARS:
            try:
                txt, engine = _ocr_image(pdf[i].render(scale=2.5).to_pil())
            except Exception as exc:
                logger.warning("OCR failed on %s page %d: %s", filename, i + 1, exc)
                continue
            if txt.strip():
                pages[i + 1] = txt
                ocr_used.add(engine)
    if ocr_used:
        method += " + OCR[" + ",".join(sorted(ocr_used)) + "]"
    return pages, method


def _docx_pages(data: bytes) -> Tuple[Dict[int, str], str]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(data))
    pages: Dict[int, list] = {1: []}
    current = 1
    saw_marker = False

    def para_text(p) -> str:
        nonlocal current, saw_marker
        out = []
        for node in p._p.iter():
            if node.tag == qn("w:lastRenderedPageBreak") or (
                    node.tag == qn("w:br") and node.get(qn("w:type")) == "page"):
                saw_marker = True
                current += 1
                pages.setdefault(current, [])
                out.append("\n@@PAGEBREAK@@\n")
            elif node.tag == qn("w:t") and node.text:
                out.append(node.text)
        return "".join(out)

    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = para_text(Paragraph(child, doc))
            _emit(pages, t, current)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                seen, cells = set(), []
                for c in row.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(" ".join(para_text(p) for p in c.paragraphs).strip())
                _emit(pages, " | ".join(x for x in cells if x), current)
    joined = {k: "\n".join(v) for k, v in pages.items() if any(s.strip() for s in v)}
    method = "docx-walk (Word page breaks)" if saw_marker else "docx-walk (no page markers in file: all text is page 1)"
    return joined, method


def _emit(pages: Dict[int, list], text: str, current: int) -> None:
    # a paragraph containing breaks is split so each piece lands on the right page
    page = current - text.count("@@PAGEBREAK@@")
    for piece in text.split("\n@@PAGEBREAK@@\n"):
        if piece.strip():
            pages.setdefault(page, []).append(piece)
        page += 1


def load_source(data: bytes, filename: str) -> SourceDoc:
    """Bytes + filename -> page-aware SourceDoc. Raises ValueError if no text can be read."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        pages, method = _pdf_pages(data, filename)
    elif ext == ".docx":
        pages, method = _docx_pages(data)
        if not pages:
            raise ValueError("DOCX contained no text")
    elif ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        from PIL import Image
        txt, engine = _ocr_image(Image.open(io.BytesIO(data)))
        pages, method = {1: txt}, f"OCR[{engine}]"
    else:
        tmp = tempfile.NamedTemporaryFile(suffix=ext or ".txt", delete=False)
        try:
            tmp.write(data)
            tmp.close()
            text = (_converter().convert(tmp.name).text_content or "")
        finally:
            os.unlink(tmp.name)
        pages, method = {1: text}, "markitdown (single page)"
    if not any(t.strip() for t in pages.values()):
        raise ValueError("No text could be extracted (empty or unreadable document)")
    return SourceDoc(filename=filename, pages=pages, method=method).build()
