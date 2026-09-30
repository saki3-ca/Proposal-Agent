"""Deterministic accuracy tests for the ToR Analyzer (no network, fake LLM)."""
import asyncio
from datetime import date

from backend.tor.analyzer import analyze, build_chunks
from backend.tor.parsers import parse_datetime, parse_duration, parse_money
from backend.tor.textutil import SourceDoc

PAGES = {
    1: "Terms of Reference\nClient: Bangladesh Example Foundation\nTender No. BEF/2026/07\n"
       "Sealed bids must be submitted by 15 October 2026 at 3:00 PM at the Dhaka office.\n"
       "Bidders shall provide a bid security of BDT 5,00,000 valid for 90 days.",
    2: "The consultant must have at least 10 years of experience in statutory audit.\n"
       "অবশ্যই বৈধ ট্রেড লাইসেন্স জমা দিতে হবে।\n"
       "A penalty of 0.5% per week shall apply for late delivery of the final report.\n"
       "The contract will be signed after approval by the Board on 1st November 2026.",
}


def _doc():
    return SourceDoc("sample.pdf", PAGES).build()


def _fake_llm(payload):
    async def llm(task, system, user, purpose=""):
        return payload, {"provider": "fake", "model": "fake-1", "error": None}
    return llm


PAYLOAD = {
    "key_facts": [
        {"field": "submission_deadline", "value": "15 October 2026 at 3:00 PM",
         "source_quote": "Sealed bids must be submitted by 15 October 2026 at 3:00 PM at the Dhaka office.", "page": 1},
        {"field": "bid_security", "value": "BDT 5,00,000",
         "source_quote": "Bidders shall provide a bid security of BDT 5,00,000 valid for 90 days.", "page": 2},  # wrong page -> corrected
        {"field": "client", "value": "Made Up Bank",  # value not in quote -> dropped
         "source_quote": "Client: Bangladesh Example Foundation", "page": 1},
        {"field": "contract_duration", "value": "24 months",  # hallucinated quote -> dropped
         "source_quote": "The contract duration is 24 months from signing.", "page": 1},
    ],
    "team": [{"position": "Consultant", "qualification_raw": "", "experience_raw": "at least 10 years",
              "source_quote": "The consultant must have at least 10 years of experience in statutory audit.", "page": 2}],
    "eligibility": [{"source_quote": "অবশ্যই বৈধ ট্রেড লাইসেন্স জমা দিতে হবে।", "page": 2}],
}


def test_hallucinated_quotes_are_dropped_and_logged():
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(PAYLOAD), today=date(2026, 9, 30)))
    dropped = {d["label"]: d["reason"] for d in snap["dropped"]}
    assert dropped["contract_duration"] == "quote not found in source"
    assert dropped["client"] == "value not found inside its own quote"
    labels = [f["label"] for f in snap["key_facts"]]
    assert "contract_duration" not in labels and "client" not in labels
    assert snap["meta"]["counts"]["dropped_quotes"] == 2


def test_page_comes_from_source_not_llm_and_values_are_parsed_by_code():
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(PAYLOAD), today=date(2026, 9, 30)))
    facts = {f["label"]: f for f in snap["key_facts"]}
    assert facts["bid_security"]["page"] == 1
    assert facts["bid_security"]["parsed"]["amount"] == "500000"
    assert facts["submission_deadline"]["parsed"] == {"date": "2026-10-15", "time": "15:00", "ambiguous": False,
                                                      "all_dates": ["2026-10-15"]}


def test_bangla_quote_verifies_and_unmapped_lists_uncovered_obligations():
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(PAYLOAD), today=date(2026, 9, 30)))
    assert snap["eligibility"] and snap["eligibility"][0]["page"] == 2
    unmapped = " ".join(u["sentence"] for u in snap["unmapped"])
    assert "penalty of 0.5%" not in unmapped          # covered by the rule-based penalty red flag
    assert "1st November 2026" in unmapped            # date sentence nobody extracted
    assert "Sealed bids must be submitted" not in unmapped  # covered by a verified item


def test_rule_based_red_flags_and_short_deadline():
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(PAYLOAD), today=date(2026, 10, 12)))
    labels = " ".join(f["label"] for f in snap["red_flags"])
    assert "penalty clause (rule-based)" in labels
    assert "3 day(s) after the analysis date" in labels


def test_conflicting_values_raise_a_flag():
    payload = {"key_facts": [
        {"field": "submission_deadline", "value": "15 October 2026", "source_quote": "Sealed bids must be submitted by 15 October 2026 at 3:00 PM", "page": 1},
        {"field": "submission_deadline", "value": "1st November 2026", "source_quote": "approval by the Board on 1st November 2026", "page": 2}]}
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(payload), today=date(2026, 9, 1)))
    assert any("Conflicting values for submission_deadline" in f["label"] for f in snap["red_flags"])


def test_no_llm_output_means_no_items():
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm({}), today=date(2026, 9, 1)))
    assert not snap["key_facts"] and snap["meta"]["counts"]["proposed_by_llm"] == 0


def test_whitespace_and_line_wrap_do_not_break_verification():
    doc = SourceDoc("w.pdf", {1: "The bidder must submit  three\ncopies of the technical proposal to the office."}).build()
    assert doc.locate("The bidder must submit three copies of the technical proposal", 1)


def test_chunks_overlap_by_a_page():
    doc = SourceDoc("c.pdf", {i: ("word " * 300) for i in range(1, 8)}).build()
    chunks = build_chunks(doc, max_chars=3500)
    assert len(chunks) > 1
    assert chunks[0]["pages"][1] == chunks[1]["pages"][0]


def test_parsers():
    assert parse_datetime("৩০ সেপ্টেম্বর ২০২৬, বিকাল ৩:০০")["date"] == "2026-09-30"
    assert parse_datetime("05/10/2026")["ambiguous"] is True
    assert parse_datetime("25/10/2026")["ambiguous"] is False
    assert parse_money("Tk. 2 lakh")["amount"] == "200000"
    assert parse_money("2% of the bid value")["percent"] == "2"
    assert parse_duration("120 (one hundred twenty) days")["unit"] == "day"


def test_exports_open_and_contain_page_numbers():
    import openpyxl
    from docx import Document
    from backend.tor.export import to_docx, to_xlsx
    snap = asyncio.run(analyze(_doc(), llm=_fake_llm(PAYLOAD), today=date(2026, 9, 30)))
    wb = openpyxl.load_workbook(to_xlsx(snap))
    assert "Key facts" in wb.sheetnames and "Unmapped - review" in wb.sheetnames
    assert wb["Key facts"]["D2"].value == 1  # page column
    text = "\n".join(p.text for p in Document(to_docx(snap)).paragraphs)
    assert "(p.2)" in text
